"""In-app file preview for the Files panel. Returns safe HTML; no network, no new services."""
import base64
import csv
import html
import io
import mimetypes
import os

MAX_BYTES = 15 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
MAX_PDF_PAGES = 12
MAX_ROWS = 300
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg"}
TEXT_EXT = {".txt", ".md", ".markdown", ".json", ".log", ".py", ".js", ".ts", ".css", ".yaml", ".yml",
            ".xml", ".ini", ".toml", ".sh", ".sql", ".rtf", ".tex"}


def safe_workspace_path(path, workspace_dir):
    """Return the real path only when it is a regular file inside workspace_dir, else None."""
    if not path:
        return None
    try:
        real = os.path.realpath(path)
        root = os.path.realpath(workspace_dir)
        if os.path.commonpath([real, root]) != root or not os.path.isfile(real):
            return None
        return real
    except (ValueError, OSError):
        return None


def _note(text):
    return f'<div class="fp-note">{html.escape(text)}</div>'


def _wrap(body):
    return f'<div class="fp-body">{body}</div>'


def _pre(text):
    clipped = len(text) > MAX_TEXT_CHARS
    out = f'<pre class="fp-text">{html.escape(text[:MAX_TEXT_CHARS])}</pre>'
    return out + (_note("Preview cut short. Download for the full file.") if clipped else "")


def _docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    parts = []

    def para_html(p):
        text = "".join(
            ("<b>" if r.bold else "") + ("<i>" if r.italic else "") + ("<u>" if r.underline else "")
            + html.escape(r.text)
            + ("</u>" if r.underline else "") + ("</i>" if r.italic else "") + ("</b>" if r.bold else "")
            for r in p.runs)
        if not text.strip():
            return ""
        style = (p.style.name if p.style is not None else "") or ""
        if style.startswith("Heading") or style == "Title":
            lvl = "".join(c for c in style if c.isdigit()) or "1"
            return f"<h{min(int(lvl) + 1, 6)}>{text}</h{min(int(lvl) + 1, 6)}>"
        if "List" in style:
            return f"<p>&bull; {text}</p>"
        align = {1: "center", 2: "right"}.get(p.alignment)
        return f'<p style="text-align:{align}">{text}</p>' if align else f"<p>{text}</p>"

    def table_html(t):
        rows = []
        for row in t.rows[:MAX_ROWS]:
            cells = "".join(f"<td>{html.escape(c.text)}</td>" for c in row.cells)
            rows.append(f"<tr>{cells}</tr>")
        return f'<table class="fp-table">{"".join(rows)}</table>'

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            parts.append(para_html(Paragraph(child, d)))
        elif tag == "tbl":
            parts.append(table_html(Table(child, d)))
    body = "".join(parts) or _note("This document has no text.")
    return _wrap(f'<div class="fp-doc">{body}</div>')


def _pdf(path):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(path)
    total = len(pdf)
    imgs = []
    for i in range(min(total, MAX_PDF_PAGES)):
        bitmap = pdf[i].render(scale=1.4)
        buf = io.BytesIO()
        bitmap.to_pil().convert("RGB").save(buf, "JPEG", quality=80)
        b64 = base64.b64encode(buf.getvalue()).decode()
        imgs.append(f'<img class="fp-page" alt="Page {i + 1}" src="data:image/jpeg;base64,{b64}">')
    extra = _note(f"Showing {MAX_PDF_PAGES} of {total} pages. Download for the rest.") if total > MAX_PDF_PAGES else ""
    return _wrap("".join(imgs) + extra)


def _image(path, ext):
    mime = mimetypes.guess_type(path)[0] or "image/png"
    b64 = base64.b64encode(open(path, "rb").read()).decode()
    return _wrap(f'<img class="fp-image" alt="{html.escape(os.path.basename(path))}" src="data:{mime};base64,{b64}">')


def _webpage(path):
    raw = open(path, "r", encoding="utf-8", errors="replace").read(MAX_TEXT_CHARS * 5)
    # Sandboxed iframe: scripts may run, but no same-origin access to this app or its cookies.
    frame = (f'<iframe class="fp-frame" sandbox="allow-scripts allow-forms allow-popups" '
             f'srcdoc="{html.escape(raw, quote=True)}"></iframe>')
    return frame + _note("Web page preview. Scripts run in a sandbox; Download to open it on its own.")


def _csv(path, delimiter=","):
    rows = []
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        for i, row in enumerate(csv.reader(f, delimiter=delimiter)):
            if i >= MAX_ROWS:
                break
            rows.append(row)
    if not rows:
        return _note("This file is empty.")
    head, rest = rows[0], rows[1:]
    th = "".join(f"<th>{html.escape(c)}</th>" for c in head)
    tr = "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>" for r in rest)
    out = f'<div class="fp-scroll"><table class="fp-table"><thead><tr>{th}</tr></thead><tbody>{tr}</tbody></table></div>'
    if len(rows) >= MAX_ROWS:
        out += _note(f"Showing the first {MAX_ROWS} rows.")
    return _wrap(out)


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets[:3]:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                break
            rows.append(["" if c is None else str(c) for c in row])
        body = "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>" for r in rows)
        out.append(f"<h4>{html.escape(ws.title)}</h4><div class=\"fp-scroll\"><table class=\"fp-table\">{body}</table></div>")
    return _wrap("".join(out) or _note("This workbook is empty."))


def _pptx(path):
    from pptx import Presentation
    out = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        out.append(f"<h4>Slide {n}</h4>" + "".join(f"<p>{html.escape(t)}</p>" for t in texts))
    return _wrap("".join(out) or _note("No text found in this deck."))


def render_preview(path):
    """HTML preview for one file. Never raises; unsupported types get a friendly note."""
    try:
        if not path or not os.path.isfile(path):
            return _note("Pick a file to preview.")
        if os.path.getsize(path) > MAX_BYTES:
            return _note("This file is too large to preview. Use Download.")
        ext = os.path.splitext(path)[1].lower()
        if ext == ".docx":
            return _docx(path)
        if ext == ".pdf":
            return _pdf(path)
        if ext in IMAGE_EXT:
            return _image(path, ext)
        if ext in (".html", ".htm"):
            return _webpage(path)
        if ext == ".csv":
            return _csv(path)
        if ext == ".tsv":
            return _csv(path, "\t")
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(path)
        if ext == ".pptx":
            return _pptx(path)
        if ext in TEXT_EXT or (mimetypes.guess_type(path)[0] or "").startswith("text/"):
            return _wrap(_pre(open(path, "r", encoding="utf-8", errors="replace").read(MAX_TEXT_CHARS + 1)))
        return _note(f"No preview for {ext or 'this'} files. Use Download.")
    except Exception as e:  # noqa: BLE001
        return _note(f"Could not preview this file ({type(e).__name__}). Use Download.")
