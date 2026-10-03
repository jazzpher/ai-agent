"""Well-formed one-page letter .docx builder, used by the create_letter_docx tool.

Small models write fragile python-docx code. This builds the layout for them:
aligned letterhead with a detailed seal (generated fictional seal or uploaded), real date,
recipient block, subject, body paragraphs, closing and signatory.
"""
import os
import tempfile
from datetime import datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x3A, 0x6B)
RED = RGBColor(0xB0, 0x1E, 0x1E)


def today_text(tz: str = "Asia/Manila") -> str:
    now = datetime.now(ZoneInfo(tz))
    return f"{now.strftime('%B')} {now.day}, {now.year}"


def resolve_date_text(date_text=None, tz="Asia/Manila"):
    """Use today's date unless a plausible current/future date was given.

    The model does not know today's date and often invents an old one, so a
    date more than 30 days in the past is replaced by today's real date.
    """
    if not date_text or not str(date_text).strip():
        return today_text()
    text = str(date_text).strip()
    today = datetime.now(ZoneInfo(tz)).date()
    for fmt in ("%B %d, %Y", "%b %d, %Y", "%Y-%m-%d", "%d %B %Y", "%B %d %Y"):
        try:
            parsed = datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        return text if (today - parsed).days <= 30 else today_text()
    return text


def _arc_text(img, text, cx, cy, radius, font, fill, start_deg, end_deg, bottom=False):
    """Place text along an arc; letters upright relative to the circle."""
    import math
    from PIL import Image, ImageDraw
    n = len(text)
    if n == 0:
        return
    for i, ch in enumerate(text):
        t = i / max(n - 1, 1)
        ang = start_deg + (end_deg - start_deg) * t  # degrees, 0=top, clockwise
        if bottom:
            ang = 180 - ang  # left-to-right along the bottom
        a = math.radians(ang)
        x, y = cx + radius * math.sin(a), cy - radius * math.cos(a)
        bb = font.getbbox(ch)
        w, h = bb[2] - bb[0] + 20, bb[3] - bb[1] + 40
        tile = Image.new("RGBA", (w * 2, h * 2), (0, 0, 0, 0))
        ImageDraw.Draw(tile).text((w, h), ch, font=font, fill=fill, anchor="mm")
        rot = tile.rotate(-(ang - 180 if bottom else ang), resample=Image.BICUBIC, expand=True)
        img.alpha_composite(rot, (int(x - rot.width / 2), int(y - rot.height / 2)))


def make_seal_png(path: str, lines=("SAMPLE", "SEAL"), size: int = 1600) -> str:
    """Draw a detailed fictional government-style seal (no real seal is imitated).

    Rope-style outer ring, text on an arc, laurel branches, mountain/sun/sea
    emblem and a star. Rendered 2x and downsampled for smooth edges."""
    import math
    from PIL import Image, ImageDraw, ImageFont
    S = size * 2
    c = S // 2
    img = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    navy, gold, cream = (24, 46, 96, 255), (184, 142, 48, 255), (250, 245, 230, 255)
    d.ellipse((4, 4, S - 4, S - 4), fill=gold)
    d.ellipse((S * 0.025, S * 0.025, S * 0.975, S * 0.975), fill=navy)
    # rope / beaded edge
    for k in range(120):
        a = 2 * math.pi * k / 120
        r = S * 0.485
        x, y = c + r * math.cos(a), c + r * math.sin(a)
        d.ellipse((x - S * 0.008, y - S * 0.008, x + S * 0.008, y + S * 0.008), fill=gold)
    d.ellipse((S * 0.06, S * 0.06, S * 0.94, S * 0.94), outline=gold, width=int(S * 0.006))
    d.ellipse((S * 0.255, S * 0.255, S * 0.745, S * 0.745), fill=cream, outline=gold, width=int(S * 0.012))
    d.ellipse((S * 0.28, S * 0.28, S * 0.72, S * 0.72), outline=navy, width=int(S * 0.004))
    words = [str(x).upper() for x in lines if str(x).strip()][:3] or ["SAMPLE"]
    top = " ".join(words) if len(words) > 1 else words[0]
    top_text = f"* {top} *"
    try:
        f1 = ImageFont.truetype("DejaVuSerif-Bold.ttf", int(S * 0.058))
        f2 = ImageFont.truetype("DejaVuSerif-Bold.ttf", int(S * 0.042))
    except OSError:
        f1 = f2 = ImageFont.load_default()
    span = min(100, 4.2 * len(top_text))
    _arc_text(img, top_text, c, c, S * 0.365, f1, cream, -span, span)
    # laurel branches hugging the emblem
    for side in (-1, 1):
        for k in range(9):
            ang = 180 + side * (14 + k * 10)
            r = S * 0.228
            a = math.radians(ang)
            x, y = c + r * math.sin(a), c - r * math.cos(a)
            leaf = Image.new("RGBA", (int(S * 0.05), int(S * 0.02)), (0, 0, 0, 0))
            ImageDraw.Draw(leaf).ellipse((0, 0, leaf.width - 1, leaf.height - 1), fill=(60, 110, 70, 255))
            rot = leaf.rotate(ang + side * 25, expand=True, resample=Image.BICUBIC)
            img.alpha_composite(rot, (int(x - rot.width / 2), int(y - rot.height / 2)))
    # emblem: sky, sun rays, mountains, sea
    em = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    e = ImageDraw.Draw(em)
    e.ellipse((S * 0.30, S * 0.30, S * 0.70, S * 0.70), fill=(214, 232, 245, 255))
    sx, sy = c, c + S * 0.02
    for k in range(16):
        a = math.pi * (k / 15)
        e.line((sx, sy, sx - math.cos(a) * S * 0.2, sy - math.sin(a) * S * 0.2), fill=(240, 196, 80, 255), width=int(S * 0.006))
    e.ellipse((sx - S * 0.05, sy - S * 0.05, sx + S * 0.05, sy + S * 0.05), fill=(232, 168, 40, 255))
    e.polygon([(c - S * 0.21, c + S * 0.12), (c - S * 0.09, c - S * 0.04), (c + S * 0.0, c + S * 0.08),
               (c + S * 0.10, c - S * 0.08), (c + S * 0.21, c + S * 0.12)], fill=(70, 100, 82, 255))
    e.polygon([(c - S * 0.21, c + S * 0.12), (c - S * 0.09, c - S * 0.04), (c - S * 0.05, c + S * 0.02)], fill=(96, 130, 104, 255))
    for k in range(4):
        y0 = c + S * (0.115 + k * 0.022)
        e.rectangle((c - S * 0.21, y0, c + S * 0.21, y0 + S * 0.016), fill=(36, 82, 140, 255) if k % 2 == 0 else (60, 120, 176, 255))
    mask = Image.new("L", (S, S), 0)
    ImageDraw.Draw(mask).ellipse((S * 0.30, S * 0.30, S * 0.70, S * 0.70), fill=255)
    img.paste(Image.composite(em, Image.new("RGBA", (S, S), (0, 0, 0, 0)), mask), (0, 0), mask)
    d.ellipse((S * 0.30, S * 0.30, S * 0.70, S * 0.70), outline=navy, width=int(S * 0.006))
    out = img.resize((size, size), Image.LANCZOS)
    out.save(path, dpi=(300, 300))
    return path


def _bottom_border(paragraph, color="1F3A6B"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    ppr.append(borders)


def _no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def build_letter(out_path, *, agency, office="", seal_lines=("SAMPLE", "SEAL"), seal_image=None,
                 date_text=None, recipient_lines=(), subject="", paragraphs=(), closing="Respectfully,",
                 signatory="", signatory_title="", footer_note="", font="Times New Roman"):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.8)
    base = doc.styles["Normal"]
    base.font.name = font
    base.font.size = Pt(11.5)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    base.paragraph_format.space_after = Pt(0)

    # Letterhead: seal left, centred agency text right, rule underneath.
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _no_borders(table)
    left, right = table.rows[0].cells
    left.width, right.width = Inches(1.3), Inches(5.2)
    tmp = None
    img = seal_image if seal_image and os.path.isfile(seal_image) else None
    if img is None:
        fd, tmp = tempfile.mkstemp(suffix=".png")
        os.close(fd)
        img = make_seal_png(tmp, seal_lines)
    left.paragraphs[0].add_run().add_picture(img, width=Inches(1.15))
    if tmp:
        os.remove(tmp)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(agency.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(17), NAVY
    if office:
        p2 = right.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(office)
        r2.font.size, r2.font.color.rgb = Pt(12), NAVY
    rule = doc.add_paragraph()
    _bottom_border(rule)
    rule.paragraph_format.space_after = Pt(14)

    def para(text="", *, bold=False, italic=False, after=0, align=None, indent=False, size=None, color=None, underline=False):
        q = doc.add_paragraph()
        if align:
            q.alignment = align
        q.paragraph_format.space_after = Pt(after)
        if indent:
            q.paragraph_format.first_line_indent = Inches(0.5)
        if text:
            run = q.add_run(text)
            run.bold, run.italic, run.underline = bold, italic, underline
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
        return q

    para(resolve_date_text(date_text), after=12)
    for i, line in enumerate(recipient_lines):
        para(line, after=12 if i == len(recipient_lines) - 1 else 0)
    if subject:
        para(subject, bold=True, underline=True, after=10)
    for text in paragraphs:
        para(text, after=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)
    para(closing, after=34)
    if signatory:
        para(signatory, bold=True)
    if signatory_title:
        para(signatory_title)
    if footer_note:
        para("", after=18)
        para(footer_note, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, color=RED)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path
