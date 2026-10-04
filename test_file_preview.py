"""File preview: pure rendering, no network."""
import base64, os, tempfile, unittest
import docx
from file_preview import render_preview, safe_workspace_path
import app
from ui_design import CSS


class RenderPreviewTest(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()

    def w(self, name, data, mode="w"):
        p = os.path.join(self.d, name)
        with open(p, mode) as f:
            f.write(data)
        return p

    def test_docx_text_headings_and_tables(self):
        doc = docx.Document()
        doc.add_heading("Liham", 1)
        doc.add_paragraph("Mahal kong <b>Ma'am</b>")
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "A"; t.rows[0].cells[1].text = "B"
        p = os.path.join(self.d, "x.docx"); doc.save(p)
        out = render_preview(p)
        self.assertIn("Liham", out); self.assertIn("&lt;b&gt;Ma&#x27;am", out)
        self.assertIn("<td>A</td>", out); self.assertNotIn("<b>Ma", out)

    def test_pdf_renders_pages_as_images(self):
        from PIL import Image
        p = os.path.join(self.d, "a.pdf"); Image.new("RGB", (200, 200), "white").save(p, "PDF")
        self.assertIn("data:image/jpeg;base64", render_preview(p))

    def test_html_goes_in_sandboxed_iframe(self):
        p = self.w("site.html", "<h1>Hi</h1><script>alert(1)</script>")
        out = render_preview(p)
        self.assertIn('sandbox="allow-scripts', out); self.assertNotIn("allow-same-origin", out)
        self.assertIn("&lt;h1&gt;Hi", out); self.assertNotIn("<script>", out)

    def test_text_csv_image_and_unknown(self):
        self.assertIn("&lt;x&gt;", render_preview(self.w("n.md", "# <x>")))
        csv_out = render_preview(self.w("t.csv", "a,b\n1,2\n"))
        self.assertIn("<th>a</th>", csv_out); self.assertIn("<td>2</td>", csv_out)
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
        self.assertIn("data:image/png;base64", render_preview(self.w("i.png", png, "wb")))
        self.assertIn("Use Download", render_preview(self.w("z.bin", b"\x00\x01", "wb")))
        self.assertIn("Pick a file", render_preview(None))

    def test_corrupt_file_does_not_raise(self):
        self.assertIn("Could not preview", render_preview(self.w("bad.docx", "not a docx")))

    def test_path_must_stay_inside_workspace(self):
        p = self.w("ok.txt", "x")
        self.assertEqual(safe_workspace_path(p, self.d), os.path.realpath(p))
        self.assertIsNone(safe_workspace_path("/etc/passwd", self.d))
        self.assertIsNone(safe_workspace_path(os.path.join(self.d, "..", "etc"), self.d))
        self.assertIsNone(safe_workspace_path(None, self.d))


class FilesPanelWiringTest(unittest.TestCase):
    def test_button_and_panel_exist_and_open_lists_files(self):
        ids = {c["props"].get("elem_id") for c in app.build_app().get_config_file()["components"]}
        for n in ("files-button", "files-panel", "files-select", "files-download", "files-preview", "files-close"):
            self.assertIn(n, ids)
        from unittest.mock import patch
        d = tempfile.mkdtemp(); p = os.path.join(d, "n.txt"); open(p, "w").write("hola")
        with patch.object(app, "WORKSPACE_DIR", d):
            panel, dd, html_out, dl = app.open_files_panel(None)
            self.assertTrue(panel["visible"]); self.assertEqual(dd["value"], p)
            self.assertIn("hola", html_out); self.assertEqual(dl["value"], os.path.realpath(p))
            self.assertIn("Pick a file", app.preview_selected("/etc/passwd")[0])

    def test_css_has_overlay_and_keeps_mobile_layout(self):
        self.assertIn("#files-panel", CSS); self.assertIn("calc(100dvh - 156px)", CSS)
